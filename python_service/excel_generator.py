from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import os

class ExcelGenerator:
    def __init__(self):
        self.export_folder = 'exports'

    def generate_excel(self, receipts_data):
        """
        Генерация Excel файла из группированных данных квитанций
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Квитанции"

        # Заголовки для группированных данных
        headers = ['ФИО', 'Группа', 'Оплачено', 'Требуется', '% оплаты', 'Статус', 'Кол-во квитанций', 'Тип банка', 'Наименование']
        ws.append(headers)

        # Стилизация заголовков
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Проверяем, является ли данные группированными (имеют поле receipts)
        is_grouped = receipts_data and len(receipts_data) > 0 and 'receipts' in receipts_data[0]

        if is_grouped:
            # Группированные данные
            for group in receipts_data:
                receipts = group.get('receipts', [])
                bank_type = receipts[0].get('bank_type', '') if receipts else ''
                university_name = receipts[0].get('university_name', '') if receipts else ''
                
                # Определяем статус на русском
                status_map = {
                    'paid': 'Оплачено',
                    'partial': 'Частично',
                    'unpaid': 'Не оплачено',
                    'overpaid': 'Переплата',
                    'pending': 'Обработка'
                }
                status = status_map.get(group.get('payment_status', ''), group.get('payment_status', ''))
                
                row = [
                    group.get('full_name', ''),
                    group.get('group', ''),
                    group.get('total_paid', 0),
                    group.get('required_amount', 0),
                    round(group.get('payment_percent', 0), 2),
                    status,
                    len(receipts),
                    bank_type,
                    university_name
                ]
                ws.append(row)
        else:
            # Старый формат (для обратной совместимости)
            for receipt in receipts_data:
                row = [
                    receipt.get('full_name', ''),
                    receipt.get('group', ''),
                    receipt.get('amount', 0),
                    receipt.get('required_amount', 0),
                    '',
                    receipt.get('status', ''),
                    1,
                    receipt.get('bank_type', ''),
                    receipt.get('university_name', '')
                ]
                ws.append(row)

        # Автоматическая ширина столбцов
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Сохранение файла
        filename = f"receipts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        wb.save(filepath)

        return filepath

    def generate_group_report(self, report_data):
        """
        Генерация Excel файла отчета по группе
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчет по группе"

        group_name = report_data.get('group', '')
        required_amount = report_data.get('required_amount', 0)
        students = report_data.get('students', [])

        # Заголовок отчета
        ws.append([f'Отчет по группе: {group_name}'])
        ws.append([f'Требуемая сумма: {required_amount} ₸'])
        ws.append([])  # Пустая строка

        # Заголовки таблицы
        headers = [
            'ФИО', 
            'Оплачено (квитанции)', 
            'Примененная оплата', 
            'Итого оплачено', 
            'Требуется', 
            'Остаток долга',
            '% оплаты', 
            'Статус',
            'Кол-во квитанций',
            'Дата применения',
            'Применил',
            'Примечания'
        ]
        ws.append(headers)

        # Стилизация заголовков
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[4]:  # Заголовки в 4-й строке
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Заполнение данных
        for student in students:
            applied_date = ''
            if student.get('applied_date'):
                try:
                    from datetime import datetime
                    date_obj = datetime.fromisoformat(student['applied_date'].replace('Z', '+00:00'))
                    applied_date = date_obj.strftime('%d.%m.%Y %H:%M')
                except:
                    applied_date = str(student.get('applied_date', ''))

            row = [
                student.get('full_name', ''),
                student.get('total_paid', 0),
                student.get('applied_amount', 0),
                student.get('total_amount', 0),
                student.get('required_amount', 0),
                student.get('remaining_debt', 0),
                round(student.get('payment_percent', 0), 2),
                student.get('payment_status', ''),
                student.get('receipt_count', 0),
                applied_date,
                student.get('applied_by', ''),
                student.get('notes', '')
            ]
            ws.append(row)

        # Автоматическая ширина столбцов
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Сохранение файла
        filename = f"group_report_{group_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        wb.save(filepath)

        return filepath

